"""Gera planilha Excel do dashboard de reembolsos com abas de dados."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


COR_CABECALHO = "34495E"
COR_CONCLUIDO = "38A169"
COR_PROCESSO = "3182CE"
COR_ZEBRA = "F5F7FA"
COR_BORDA = "D5DDE3"
COR_TEXTO = "2C3E50"

MOEDA = '"R$" #,##0.00'
PERCENTUAL = '0.00"%"'
INTEIRO = "#,##0"

THIN = Border(
    left=Side(style="thin", color=COR_BORDA),
    right=Side(style="thin", color=COR_BORDA),
    top=Side(style="thin", color=COR_BORDA),
    bottom=Side(style="thin", color=COR_BORDA),
)


def _fill(color):
    return PatternFill("solid", fgColor=color)


def _font(bold=False, size=11, color="FFFFFF", name="Calibri"):
    return Font(name=name, bold=bold, size=size, color=color)


def _align(h="center", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def _n(value):
    try:
        if value is None or value == "":
            return 0
        return float(value)
    except (TypeError, ValueError):
        return 0


def _texto(value):
    if value is None:
        return ""
    return str(value)


def _ajustar_colunas(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _pintar_cabecalho(ws, row, cols, fill_color=COR_CABECALHO):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _fill(fill_color)
        cell.font = _font(bold=True, size=11)
        cell.alignment = _align()
        cell.border = THIN


def _tabela(ws, name, ref):
    tabela = Table(displayName=name, ref=ref)
    tabela.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tabela)


def _sheet_kpis(wb, kpis):
    ws = wb.create_sheet("KPIs")
    headers = ["Indicador", "Valor (R$)", "Quantidade", "% do total"]
    ws.append(headers)
    _pintar_cabecalho(ws, 1, 4)

    linhas = [
        ("Total geral", _n(kpis.get("total_geral_valor")), _n(kpis.get("total_geral_count")), 100),
        ("Concluídos", _n(kpis.get("total_concluido")), _n(kpis.get("count_concluido")), _n(kpis.get("taxa_aprovacao"))),
        ("Em processo", _n(kpis.get("total_em_processo")), _n(kpis.get("count_em_processo")), _n(kpis.get("taxa_em_processo"))),
        ("Rejeitados", _n(kpis.get("total_rejeitado")), _n(kpis.get("count_rejeitado")), _n(kpis.get("taxa_rejeicao"))),
    ]
    for linha in linhas:
        ws.append(linha)

    for row in range(2, 6):
        ws.cell(row=row, column=2).number_format = MOEDA
        ws.cell(row=row, column=3).number_format = INTEIRO
        ws.cell(row=row, column=4).number_format = PERCENTUAL
        for col in range(1, 5):
            ws.cell(row=row, column=col).alignment = _align(h="left" if col == 1 else "center")
            ws.cell(row=row, column=col).font = Font(name="Calibri", color=COR_TEXTO)
            if row % 2 == 0:
                ws.cell(row=row, column=col).fill = _fill(COR_ZEBRA)

    _ajustar_colunas(ws, {"A": 22, "B": 18, "C": 16, "D": 16})
    _tabela(ws, "TabelaKPIs", "A1:D5")
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = COR_CABECALHO
    return ws


def _sheet_areas(wb, areas):
    ws = wb.create_sheet("Por Área")
    headers = [
        "Área",
        "Concluído (R$)",
        "Em processo (R$)",
        "Rejeitado (R$)",
        "Total (R$)",
        "Qtd. concluídos",
        "Qtd. em processo",
        "Qtd. rejeitados",
        "Qtd. total",
        "Taxa aprovação (%)",
        "Taxa rejeição (%)",
        "Taxa em processo (%)",
    ]
    ws.append(headers)
    _pintar_cabecalho(ws, 1, len(headers))

    for area in areas:
        ws.append([
            _texto(area.get("centro_label") or area.get("centro")),
            _n(area.get("concluido_total")),
            _n(area.get("em_processo_total")),
            _n(area.get("rejeitado_total")),
            _n(area.get("total_geral")),
            _n(area.get("concluido_count")),
            _n(area.get("em_processo_count")),
            _n(area.get("rejeitado_count")),
            _n(area.get("count_geral")),
            _n(area.get("taxa_aprovacao")),
            _n(area.get("taxa_rejeicao")),
            _n(area.get("taxa_em_processo")),
        ])

    last = max(len(areas) + 1, 2)
    for row in range(2, last + 1):
        for col in (2, 3, 4, 5):
            ws.cell(row=row, column=col).number_format = MOEDA
        for col in (6, 7, 8, 9):
            ws.cell(row=row, column=col).number_format = INTEIRO
        for col in (10, 11, 12):
            ws.cell(row=row, column=col).number_format = PERCENTUAL
        for col in range(1, 13):
            ws.cell(row=row, column=col).alignment = _align(h="left" if col == 1 else "center")
            ws.cell(row=row, column=col).font = Font(name="Calibri", color=COR_TEXTO)
            if row % 2 == 0:
                ws.cell(row=row, column=col).fill = _fill(COR_ZEBRA)

    _ajustar_colunas(ws, {get_column_letter(i): 18 for i in range(1, 13)})
    ws.column_dimensions["A"].width = 28
    if areas:
        _tabela(ws, "TabelaPorArea", f"A1:L{last}")
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = COR_CONCLUIDO
    return ws


def _sheet_tipos(wb, tipos):
    ws = wb.create_sheet("Por Tipo")
    ws.append(["Tipo de despesa", "Valor (R$)"])
    _pintar_cabecalho(ws, 1, 2, COR_PROCESSO)

    labels = tipos.get("labels") or []
    valores = tipos.get("valores") or []
    for label, valor in zip(labels, valores):
        ws.append([_texto(label), _n(valor)])

    last = max(len(labels) + 1, 2)
    for row in range(2, last + 1):
        ws.cell(row=row, column=2).number_format = MOEDA
        ws.cell(row=row, column=1).font = Font(name="Calibri", color=COR_TEXTO)
        ws.cell(row=row, column=2).font = Font(name="Calibri", color=COR_TEXTO)
        if row % 2 == 0:
            ws.cell(row=row, column=1).fill = _fill(COR_ZEBRA)
            ws.cell(row=row, column=2).fill = _fill(COR_ZEBRA)

    _ajustar_colunas(ws, {"A": 36, "B": 18})
    if labels:
        _tabela(ws, "TabelaPorTipo", f"A1:B{last}")
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = COR_PROCESSO
    return ws


def _sheet_meses(wb, meses):
    ws = wb.create_sheet("Evolução Mensal")
    ws.append(["Mês", "Concluído (R$)", "Em processo (R$)", "Rejeitado (R$)"])
    _pintar_cabecalho(ws, 1, 4)

    labels = meses.get("labels") or []
    concluidos = meses.get("concluidos") or []
    em_processo = meses.get("em_processo") or []
    rejeitados = meses.get("rejeitados") or []
    for i, label in enumerate(labels):
        ws.append([
            _texto(label),
            _n(concluidos[i] if i < len(concluidos) else 0),
            _n(em_processo[i] if i < len(em_processo) else 0),
            _n(rejeitados[i] if i < len(rejeitados) else 0),
        ])

    last = max(len(labels) + 1, 2)
    for row in range(2, last + 1):
        for col in range(2, 5):
            ws.cell(row=row, column=col).number_format = MOEDA
            ws.cell(row=row, column=col).alignment = _align()
        ws.cell(row=row, column=1).font = Font(name="Calibri", color=COR_TEXTO)
        if row % 2 == 0:
            for col in range(1, 5):
                ws.cell(row=row, column=col).fill = _fill(COR_ZEBRA)

    _ajustar_colunas(ws, {"A": 22, "B": 20, "C": 22, "D": 20})
    if labels:
        _tabela(ws, "TabelaEvolucaoMensal", f"A1:D{last}")
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "F59E0B"
    return ws


def _sheet_tabela(wb, titulo, headers, rows, table_name, tab_color, money_cols=None):
    ws = wb.create_sheet(titulo)
    ws.append(headers)
    _pintar_cabecalho(ws, 1, len(headers), tab_color)
    money_cols = set(money_cols or [])

    for row in rows:
        ws.append(list(row))

    last = max(len(rows) + 1, 2)
    for r in range(2, last + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = _align(h="left" if c == 1 else "center", wrap=False)
            cell.font = Font(name="Calibri", color=COR_TEXTO, size=10)
            if c in money_cols:
                cell.number_format = MOEDA
            if r % 2 == 0:
                cell.fill = _fill(COR_ZEBRA)

    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(str(header)) + 4, 14), 36)

    if rows:
        last_col = get_column_letter(len(headers))
        _tabela(ws, table_name, f"A1:{last_col}{last}")
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = tab_color
    return ws


def gerar_excel_dashboard(
    *,
    kpis,
    filtros,
    insights,
    areas,
    tipos,
    meses,
    solicitacoes_rows,
    itens_rows,
    gerado_em,
):
    wb = Workbook()
    ws_padrao = wb.active

    ws_kpis = _sheet_kpis(wb, kpis)
    _sheet_areas(wb, areas)
    _sheet_tipos(wb, tipos)
    _sheet_meses(wb, meses)

    headers_sol = [
        "ID", "Solicitante", "Centro de custo", "Status", "Status descritivo",
        "Valor total (R$)", "Pago", "Concluído", "Data da solicitação",
        "Data conclusão", "Forma de pagamento",
    ]
    _sheet_tabela(
        wb, "Solicitações", headers_sol, solicitacoes_rows,
        "TabelaSolicitacoes", "34495E", money_cols={6},
    )

    headers_itens = [
        "ID solicitação", "Solicitante", "Centro de custo", "Status",
        "Tipo de despesa", "Código orçamento", "Data da despesa",
        "Descrição", "Valor (R$)", "KM",
    ]
    _sheet_tabela(
        wb, "Itens", headers_itens, itens_rows,
        "TabelaItens", "8B5CF6", money_cols={9},
    )

    wb.remove(ws_padrao)
    wb.active = ws_kpis
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out
